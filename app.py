import streamlit as st
import pandas as pd

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Alignment

st.set_page_config(
    page_title="構造出欠管理",
    layout="wide"
)

members = pd.read_excel(
    "members.xlsx"
)

page = st.sidebar.radio(
    "メニュー",
    [
        "出欠登録",
        "Excel出力"
    ]
)

# ==================================================
# 出欠登録画面
# ==================================================

if page == "出欠登録":

    st.title(
        "構造出欠登録"
    )

    date = st.date_input(
        "日付"
    )

    search = st.text_input(
        "氏名検索（ひらがな・漢字可）"
    )

    if search:

        filtered = members[
            members["氏名"]
            .astype(str)
            .str.contains(
                search,
                na=False
            )
            |
            members["ふりがな"]
            .astype(str)
            .str.contains(
                search,
                na=False
            )
        ]

    else:

        filtered = members

    names = filtered[
        "氏名"
    ].tolist()

    if len(names) > 0:

        selected_name = st.selectbox(
            "氏名",
            names
        )

        attendance = st.radio(
            "出欠",
            [
                "出席",
                "欠席"
            ]
        )

        if st.button(
            "登録"
        ):

            new_data = pd.DataFrame(
                [
                    {
                        "日付": str(date),
                        "氏名": selected_name,
                        "出欠": attendance
                    }
                ]
            )

            try:

                old_data = pd.read_csv(
                    "attendance.csv",
                    encoding="utf-8-sig"
                )

                data = pd.concat(
                    [
                        old_data,
                        new_data
                    ],
                    ignore_index=True
                )

            except:

                data = new_data

            data.to_csv(
                "attendance.csv",
                index=False,
                encoding="utf-8-sig"
            )

            st.success(
                f"{selected_name} を {attendance} で登録しました"
            )

    else:

        st.warning(
            "該当する氏名がありません"
        )

    st.divider()

    st.subheader(
        "本日の登録一覧"
    )

    try:

        data = pd.read_csv(
            "attendance.csv",
            encoding="utf-8-sig"
        )

        today = data[
            data["日付"]
            == str(date)
        ]

        today = today.drop_duplicates(
            subset=["氏名"],
            keep="last"
        )

        st.dataframe(
            today,
            use_container_width=True
        )

    except:

        st.info(
            "まだ登録はありません"
        )


# ==================================================
# Excel出力画面
# ==================================================

elif page == "Excel出力":

    st.title(
        "構造出欠リスト出力"
    )

    target_date = st.date_input(
        "出力対象日"
    )

    try:

        data = pd.read_csv(
            "attendance.csv",
            encoding="utf-8-sig"
        )

        target = data[
            data["日付"]
            == str(target_date)
        ]

        # 同一人物は最新のみ採用
        target = target.drop_duplicates(
            subset=["氏名"],
            keep="last"
        )

        attend = target[
            target["出欠"]
            == "出席"
        ]

        registered_names = target[
            "氏名"
        ].tolist()

        all_names = members[
            "氏名"
        ].tolist()

        unanswered = [
            x for x in all_names
            if x not in registered_names
        ]

        absent_names = target[
            target["出欠"]
            == "欠席"
        ]["氏名"].tolist()

        absent_names.extend(
            unanswered
        )

        absent = pd.DataFrame(
            {
                "氏名": sorted(
                    list(
                        set(absent_names)
                    )
                )
            }
        )

        col1, col2 = st.columns(
            2
        )

        with col1:

            st.subheader(
                f"出席者 ({len(attend)}名)"
            )

            st.dataframe(
                attend[
                    ["氏名"]
                ],
                use_container_width=True
            )

        with col2:

            st.subheader(
                f"欠席者 ({len(absent)}名)"
            )

            st.dataframe(
                absent,
                use_container_width=True
            )

        yy = str(
            target_date.year
        )[2:]

        mm = (
            f"{target_date.month:02d}"
        )

        dd = (
            f"{target_date.day:02d}"
        )

        if st.button(
            "Excel作成"
        ):

            wb = Workbook()

            ws = wb.active

            ws.title = (
                "構造出欠リスト"
            )

            # タイトル
            ws.merge_cells(
                "A1:B1"
            )

            ws["A1"] = (
                "構造出欠リスト"
            )

            ws["A1"].font = Font(
                bold=True,
                size=16
            )

            ws["A1"].alignment = (
                Alignment(
                    horizontal="center"
                )
            )

            # 日付
            ws["A3"] = "対象日"

            ws["B3"] = (
                f"{target_date.year}/"
                f"{target_date.month:02d}/"
                f"{target_date.day:02d}"
            )

            # 人数
            ws["A4"] = "出席者数"
            ws["B4"] = len(attend)

            ws["A5"] = "欠席者数"
            ws["B5"] = len(absent)

            # ヘッダー
            ws["A7"] = "出席"
            ws["B7"] = "欠席"

            header_fill = (
                PatternFill(
                    fill_type="solid",
                    fgColor="D9EAF7"
                )
            )

            ws["A7"].fill = (
                header_fill
            )

            ws["B7"].fill = (
                header_fill
            )

            ws["A7"].font = Font(
                bold=True
            )

            ws["B7"].font = Font(
                bold=True
            )

            attend_names = (
                attend["氏名"]
                .tolist()
            )

            absent_names = (
                absent["氏名"]
                .tolist()
            )

            max_rows = max(
                len(attend_names),
                len(absent_names)
            )

            while len(attend_names) < max_rows:
                attend_names.append("")

            while len(absent_names) < max_rows:
                absent_names.append("")

            row = 8

            for a, b in zip(
                attend_names,
                absent_names
            ):

                ws.cell(
                    row=row,
                    column=1,
                    value=a
                )

                ws.cell(
                    row=row,
                    column=2,
                    value=b
                )

                row += 1

            ws.column_dimensions[
                "A"
            ].width = 30

            ws.column_dimensions[
                "B"
            ].width = 30

            output = BytesIO()

            wb.save(
                output
            )

            output.seek(0)

            filename = (
                f"{yy}{mm}{dd}"
                "_構造出欠リスト.xlsx"
            )

            st.download_button(
                label="Excelダウンロード",
                data=output,
                file_name=filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

    except Exception as e:

        st.error(
            str(e)
        )

